import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn
import zipfile
import tempfile
import shutil
import io
import os

def generate_report():
    orig_path = "Wegagen Bank Sc(software development)_original_backup.docx"
    orig_doc = docx.Document(orig_path)
    
    # Extract images
    imgs = {}
    for rId in ['rId8', 'rId9', 'rId10', 'rId11', 'rId12', 'rId13', 'rId14']:
        rel = orig_doc.part.rels[rId]
        imgs[rId] = rel.target_part.blob

    doc = docx.Document()
    
    # Set page margins to standard 1 inch
    section = doc.sections[0]
    section.top_margin = Inches(1.0)
    section.bottom_margin = Inches(1.0)
    section.left_margin = Inches(1.0)
    section.right_margin = Inches(1.0)
    section.page_width = Inches(8.5)
    section.page_height = Inches(11.0)
    
    # Enable different first page for prelims (cover page unnumbered)
    section.different_first_page_header_footer = True
    
    # Configure styles
    styles = doc.styles
    
    # Normal / Body
    norm = styles['Normal']
    norm.font.name = 'Times New Roman'
    norm.font.size = Pt(12)
    norm.font.color.rgb = RGBColor(0x11, 0x11, 0x11)
    norm.paragraph_format.line_spacing = 1.15
    norm.paragraph_format.space_after = Pt(4)
    norm.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    # Heading 1
    h1 = styles['Heading 1']
    h1.font.name = 'Arial'
    h1.font.size = Pt(15)
    h1.font.bold = True
    h1.font.color.rgb = RGBColor(0x00, 0x20, 0x60) # Deep Navy
    h1.paragraph_format.space_before = Pt(14)
    h1.paragraph_format.space_after = Pt(6)
    h1.paragraph_format.keep_with_next = True

    # Heading 2
    h2 = styles['Heading 2']
    h2.font.name = 'Arial'
    h2.font.size = Pt(12.5)
    h2.font.bold = True
    h2.font.color.rgb = RGBColor(0x1F, 0x4E, 0x78)
    h2.paragraph_format.space_before = Pt(10)
    h2.paragraph_format.space_after = Pt(4)
    h2.paragraph_format.keep_with_next = True

    # Heading 3
    h3 = styles['Heading 3']
    h3.font.name = 'Arial'
    h3.font.size = Pt(11)
    h3.font.bold = True
    h3.font.color.rgb = RGBColor(0x2E, 0x75, 0xB5)
    h3.paragraph_format.space_before = Pt(6)
    h3.paragraph_format.space_after = Pt(2)
    h3.paragraph_format.keep_with_next = True

    # Footer for Section 1 (Preliminaries: Roman numerals)
    f1 = section.footer.paragraphs[0]
    f1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    f1.paragraph_format.space_before = Pt(0)
    f1.paragraph_format.space_after = Pt(0)
    r_f1 = f1.add_run()
    r_f1.font.name = 'Times New Roman'
    r_f1.font.size = Pt(10)
    fld_xml = '<w:fldSimple ' + nsdecls('w') + ' w:instr="PAGE"/>'
    r_f1._r.append(parse_xml(fld_xml))
    
    pg_num_xml = '<w:pgNumType ' + nsdecls('w') + ' w:fmt="lowerRoman" w:start="1"/>'
    section._sectPr.append(parse_xml(pg_num_xml))

    # Helper functions
    def add_p(text="", style='Normal', space_before=0, space_after=4, align=WD_ALIGN_PARAGRAPH.JUSTIFY, bold=False, italic=False, font_size=12, color=None):
        p = doc.add_paragraph(style=style)
        p.paragraph_format.space_before = Pt(space_before)
        p.paragraph_format.space_after = Pt(space_after)
        p.paragraph_format.alignment = align
        if text:
            r = p.add_run(text)
            r.bold = bold
            r.italic = italic
            r.font.size = Pt(font_size)
            if color:
                r.font.color.rgb = color
        return p

    def add_bullet(text, bold_prefix=None, space_after=3):
        p = doc.add_paragraph(style='List Paragraph')
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(space_after)
        p.paragraph_format.line_spacing = 1.15
        p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
        if bold_prefix:
            r_bold = p.add_run(bold_prefix)
            r_bold.bold = True
            r_bold.font.name = 'Times New Roman'
            r_bold.font.size = Pt(12)
        r_text = p.add_run(text)
        r_text.font.name = 'Times New Roman'
        r_text.font.size = Pt(12)
        return p

    def add_caption(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.space_after = Pt(8)
        p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(text)
        r.font.name = 'Times New Roman'
        r.font.size = Pt(10)
        r.bold = True
        r.italic = True
        r.font.color.rgb = RGBColor(0x33, 0x33, 0x33)
        return p

    def add_toc_item(title, page_str, level=1):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(1)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.line_spacing = 1.15
        p.paragraph_format.tab_stops.add_tab_stop(Inches(6.5), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
        
        if level == 1:
            p.paragraph_format.left_indent = Inches(0)
            r1 = p.add_run(title)
            r1.font.name = 'Arial'
            r1.bold = True
            r1.font.size = Pt(10.5)
            r1.font.color.rgb = RGBColor(0x00, 0x20, 0x60)
            r2 = p.add_run(f'\t{page_str}')
            r2.font.name = 'Times New Roman'
            r2.bold = True
            r2.font.size = Pt(11)
        elif level == 2:
            p.paragraph_format.left_indent = Inches(0.25)
            r1 = p.add_run(title)
            r1.font.name = 'Times New Roman'
            r1.bold = False
            r1.font.size = Pt(11)
            r2 = p.add_run(f'\t{page_str}')
            r2.font.name = 'Times New Roman'
            r2.bold = False
            r2.font.size = Pt(11)
        elif level == 3:
            p.paragraph_format.left_indent = Inches(0.5)
            r1 = p.add_run(title)
            r1.font.name = 'Times New Roman'
            r1.bold = False
            r1.font.size = Pt(10.5)
            r2 = p.add_run(f'\t{page_str}')
            r2.font.name = 'Times New Roman'
            r2.bold = False
            r2.font.size = Pt(10.5)

    # ----------------------------------------------------
    # COVER PAGE
    # ----------------------------------------------------
    p_logo = doc.add_paragraph()
    p_logo.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_logo.paragraph_format.space_before = Pt(0)
    p_logo.paragraph_format.space_after = Pt(10)
    r_logo = p_logo.add_run()
    r_logo.add_picture(io.BytesIO(imgs['rId8']), width=Inches(1.75))

    add_p("HAWASSA UNIVERSITY", space_after=2, align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, font_size=15, color=RGBColor(0x00, 0x20, 0x60))
    add_p("INSTITUTE OF TECHNOLOGY", space_after=2, align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, font_size=13, color=RGBColor(0x00, 0x20, 0x60))
    add_p("FACULTY OF INFORMATICS", space_after=2, align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, font_size=12, color=RGBColor(0x1F, 0x4E, 0x78))
    add_p("DEPARTMENT OF COMPUTER SCIENCE", space_after=18, align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, font_size=12, color=RGBColor(0x1F, 0x4E, 0x78))

    add_p("Internship Report on Software Development", space_after=2, align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, font_size=14, color=RGBColor(0x00, 0x20, 0x60))
    add_p("at Wegagen Bank", space_after=24, align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, font_size=14, color=RGBColor(0x00, 0x20, 0x60))

    # Meta box
    meta_p = add_p("Student Name:  Yeamlak Sisay", space_after=3, align=WD_ALIGN_PARAGRAPH.LEFT, bold=True, font_size=11)
    add_p("Mentor Name:  Mr. Habtamu Assegahegn", space_after=3, align=WD_ALIGN_PARAGRAPH.LEFT, bold=True, font_size=11)
    add_p("Hosting Company:  Wegagen Bank S.C.", space_after=3, align=WD_ALIGN_PARAGRAPH.LEFT, bold=True, font_size=11)
    add_p("Duration of Apprenticeship:  June 23, 2026 - September 1, 2026", space_after=3, align=WD_ALIGN_PARAGRAPH.LEFT, bold=True, font_size=11)
    add_p("Host Department:  Core Application Management Service", space_after=16, align=WD_ALIGN_PARAGRAPH.LEFT, bold=True, font_size=11)
    add_p("Date of Submission:  September 18, 2026", space_after=0, align=WD_ALIGN_PARAGRAPH.RIGHT, bold=True, italic=True, font_size=11)

    # ----------------------------------------------------
    # PRELIMINARY PAGES
    # ----------------------------------------------------
    
    # 1. Declaration
    doc.add_page_break()
    add_p("Declaration", style='Heading 1', space_before=0, space_after=8, align=WD_ALIGN_PARAGRAPH.LEFT)
    add_p("I declare that this internship activity report is my original work, prepared based on the activities, experiences, and knowledge gained during my internship period from June 23 to September 1, 2026 at Wegagen Bank where I participated in software development.")
    add_p("All sources of information used in preparing this report have been appropriately acknowledged.")
    p_dec_sign = add_p("", space_before=12, space_after=4)
    r_dec1 = p_dec_sign.add_run("Student Name: Yeamlak Sisay\nStudent ID: 3866/16")
    r_dec1.bold = True

    # 2. Approval of the Mentor
    doc.add_page_break()
    add_p("Approval of the Mentor", style='Heading 1', space_before=0, space_after=8, align=WD_ALIGN_PARAGRAPH.LEFT)
    add_p('This serves as confirmation that the practical attachment report titled "Software Development Internship at Wegagen Bank" has undergone a thorough review and is officially approved for submission.')
    add_p('The review and evaluation of this report were carried out by the Mentor Habtamu Assegahegn, who confirms that the content meets the required academic and professional standards for a practical attachment report and is fit for submission.')
    p_appr_sign = add_p("", space_before=12, space_after=4)
    r_appr = p_appr_sign.add_run("Mentor Name: Mr. Habtamu Assegahegn\nSignature: ______________________\nDate: __________________________")
    r_appr.bold = True

    # 3. Acknowledgements
    doc.add_page_break()
    add_p("Acknowledgements", style='Heading 1', space_before=0, space_after=8, align=WD_ALIGN_PARAGRAPH.LEFT)
    add_p("I would like to express my sincere appreciation to Wegagen Bank for accepting me as an intern and providing me with a valuable professional learning environment. I would like to thank the staff members of the Core Application Management Service Department for their guidance, support, cooperation, and willingness to share their knowledge and experience.")
    add_p("My special thanks go to my department advisor, Mr. Habtamu Assegahegn and his teammates Mr. Abdulrezak and Mr. Robel, for their supervision, constructive feedback, and continuous guidance throughout the internship and report preparation.")

    # 4. Executive Summary
    doc.add_page_break()
    add_p("Executive Summary", style='Heading 1', space_before=0, space_after=8, align=WD_ALIGN_PARAGRAPH.LEFT)
    add_p("This report documents a two-and-a-half-month internship completed at Wegagen Bank S.C., within the Core Application Management Services department, from June 23 to September 1, 2026. The internship centered on the design and development of a mobile application — the Smart Agent & Branch Queue and Forex Finder system — aimed at addressing a practical operational problem observed at the bank's branches: long customer queues for routine service.")
    add_p("The first phase of the internship focused on orientation and study, gaining an understanding of how quality software products are delivered within a regulated banking environment and how the bank manages and supports its core systems. Building on this foundation, the internship moved into hands-on development, using Node.js and Express to build a stable backend — including routing, database connectivity, middleware, and services — before layering a React Native and Expo frontend on top, developed in TypeScript. The resulting application allows customers to reserve their position in a branch queue remotely rather than waiting in line, while also giving staff and customers visibility into daily foreign-exchange rates, in line with the bank's internal rate-update policy. Additional features included loan and card management support and appointment scheduling, with Clerk used for authentication and email-based notifications.")
    add_p("Beyond the technical build, the internship provided insight into how software decisions are shaped by organizational policy and risk considerations in a financial institution, and how backend architecture decisions made early in a project affect everything built afterward. The experience strengthened both technical skills, in backend and mobile application development, and professional skills, in working within a structured, policy-driven organization, and it clarified a personal interest in pursuing backend development further.")

    # 5. Table of Contents
    doc.add_page_break()
    p_toc_hdr = add_p("Table of Contents", style='Heading 1', space_before=0, space_after=8, align=WD_ALIGN_PARAGRAPH.LEFT)
    
    # Dynamic TOC XML Field
    fld_toc = p_toc_hdr.add_run()
    r_toc = fld_toc._r
    fldChar1 = parse_xml(r'<w:fldChar %s w:fldCharType="begin"/>' % nsdecls('w'))
    instrText = parse_xml(r'<w:instrText %s xml:space="preserve"> TOC \o "1-3" \h \z \u </w:instrText>' % nsdecls('w'))
    fldChar2 = parse_xml(r'<w:fldChar %s w:fldCharType="separate"/>' % nsdecls('w'))
    fldChar3 = parse_xml(r'<w:fldChar %s w:fldCharType="end"/>' % nsdecls('w'))
    r_toc.append(fldChar1)
    r_toc.append(instrText)
    r_toc.append(fldChar2)
    r_toc.append(fldChar3)

    # Populated TOC lines
    add_toc_item("Declaration", "ii", level=1)
    add_toc_item("Approval of the Mentor", "iii", level=1)
    add_toc_item("Acknowledgements", "iv", level=1)
    add_toc_item("Executive Summary", "v", level=1)
    add_toc_item("List of Figures", "vi", level=1)
    
    add_toc_item("CHAPTER ONE: BACKGROUND OF WEGAGEN BANK", "1", level=1)
    add_toc_item("1.1 Introduction", "1", level=2)
    add_toc_item("1.2 Brief History of Wegagen Bank", "1", level=2)
    add_toc_item("1.3 Main Products and Services of Wegagen Bank", "2", level=2)
    add_toc_item("1.3.1 Deposit and Savings Services", "2", level=3)
    add_toc_item("1.3.2 Credit and Loan Facilities", "2", level=3)
    add_toc_item("1.3.3 Digital and Electronic Banking (e-Banking)", "3", level=3)
    add_toc_item("1.3.4 Interest-Free Banking (Wegagen Amanah)", "3", level=3)
    add_toc_item("1.4 Main Customers and End Users of Wegagen Bank's Products and Services", "3", level=2)
    add_toc_item("1.4.1 Retail and Individual Banking Customers", "3", level=3)
    add_toc_item("1.4.2 Micro, Small, and Medium Enterprises (MSMEs)", "4", level=3)
    add_toc_item("1.4.3 Corporate and Commercial Entities", "4", level=3)
    add_toc_item("1.5 Organization and Structure", "5", level=2)
    add_toc_item("1.6 Workflow in Core Application Management Services", "6", level=2)

    add_toc_item("CHAPTER TWO: PRACTICAL ATTACHMENT EXPERIENCE", "7", level=1)
    add_toc_item("2.1 How I Joined the Company", "7", level=2)
    add_toc_item("2.1.1 Placement within Wegagen Bank: Core Application Management Service", "7", level=3)
    add_toc_item("2.1.2 Supervision and Technical Leadership", "8", level=3)
    add_toc_item("2.2 Workflow in the Core Application Management Service Department", "8", level=2)
    add_toc_item("2.3 Work Tasks Executed During the Practical Attachment", "9", level=2)
    add_toc_item("2.3.1 Problem Identification and Requirement Analysis", "9", level=3)
    add_toc_item("2.3.2 System Planning and Design", "9", level=3)
    add_toc_item("2.3.3 Backend Architecture and Database Implementation", "10", level=3)
    add_toc_item("2.3.4 Mobile Application Development", "11", level=3)
    add_toc_item("2.3.5 Queue Management and Foreign-Exchange Rate Implementation", "12", level=3)
    add_toc_item("2.3.6 Quality Assurance, System Documentation, and Project Delivery", "13", level=3)
    add_toc_item("2.4 Materials and Methodology", "14", level=2)
    add_toc_item("2.4.1 Development Environment and Tools", "14", level=3)
    add_toc_item("2.4.2 Frontend and Mobile Technologies", "14", level=3)
    add_toc_item("2.4.3 Backend and Database Technologies", "15", level=3)
    add_toc_item("2.4.4 Authentication, Real-Time Communication, and Testing", "15", level=3)
    add_toc_item("2.4.5 Version Control and Methodology", "16", level=3)
    add_toc_item("2.5 Evaluation of Work Performance", "17", level=2)
    add_toc_item("2.6 Challenges Faced During the Practical Attachment", "18", level=2)
    add_toc_item("2.6.1 Learning React Native and Expo from the Ground Up", "18", level=3)
    add_toc_item("2.6.2 Designing the Application Architecture from Scratch", "18", level=3)
    add_toc_item("2.6.3 API Integration and Client-Server Communication Bug", "19", level=3)
    add_toc_item("2.6.4 Balancing Learning with a Demanding Commute and Schedule", "19", level=3)
    add_toc_item("2.7 Measures Taken to Overcome Challenges", "19", level=2)
    add_toc_item("2.7.1 Overcoming the React Native and Expo Learning Curve", "19", level=3)
    add_toc_item("2.7.2 Iterative Design and Incremental Architecture", "20", level=3)
    add_toc_item("2.7.3 Systematic Debugging of API Endpoints", "20", level=3)
    add_toc_item("2.7.4 Time Management and Structured Routine", "20", level=3)

    add_toc_item("CHAPTER THREE: BENEFITS GAINED", "21", level=1)
    add_toc_item("3.1 Practical Skills", "21", level=2)
    add_toc_item("3.2 Theoretical Understanding Gained", "22", level=2)
    add_toc_item("3.2.1 Mobile Application Architecture", "22", level=3)
    add_toc_item("3.2.2 API Design and Client-Server Communication", "22", level=3)
    add_toc_item("3.2.3 Software Development in a Real-World, Production-Oriented Environment", "23", level=3)
    add_toc_item("3.3 Interpersonal and Communication Skills Gained", "23", level=2)
    add_toc_item("3.4 Teamwork Skills", "24", level=2)
    add_toc_item("3.5 Understanding of Work Ethics", "24", level=2)

    add_toc_item("CHAPTER FOUR: CONCLUSION AND RECOMMENDATIONS", "25", level=1)
    add_toc_item("4.1 Conclusion", "25", level=2)
    add_toc_item("4.2 Recommendations", "26", level=2)
    add_toc_item("4.2.1 Structured Onboarding for Interns", "26", level=3)
    add_toc_item("4.2.2 Pre-Internship Skill Preparation", "26", level=3)
    add_toc_item("4.2.3 University and Industry Collaboration", "27", level=3)

    add_toc_item("REFERENCES", "28", level=1)

    # 6. List of Figures
    doc.add_page_break()
    add_p("List of Figures", style='Heading 1', space_before=0, space_after=8, align=WD_ALIGN_PARAGRAPH.LEFT)
    add_toc_item("Figure 1.1: Wegagen Bank Headquarters", "2", level=2)
    add_toc_item("Figure 2.1: Smart Agent Mobile App – Home and Branch Discovery Screens", "11", level=2)
    add_toc_item("Figure 2.2: Dedicated Teller and Employee Queue Management Interface", "12", level=2)
    add_toc_item("Figure 2.3: Backend REST API Swagger/OpenAPI Documentation", "13", level=2)

    # ----------------------------------------------------
    # SECTION 2: MAIN CHAPTERS (Arabic page numbers 1, 2, 3...)
    # ----------------------------------------------------
    doc.add_page_break()
    s2 = doc.add_section()
    s2.top_margin = Inches(1.0)
    s2.bottom_margin = Inches(1.0)
    s2.left_margin = Inches(1.0)
    s2.right_margin = Inches(1.0)
    s2.page_width = Inches(8.5)
    s2.page_height = Inches(11.0)
    s2.different_first_page_header_footer = False
    s2.footer.is_linked_to_previous = False
    s2.header.is_linked_to_previous = False
    
    # Section 2 Footer
    f2 = s2.footer.paragraphs[0]
    f2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    f2.paragraph_format.space_before = Pt(0)
    f2.paragraph_format.space_after = Pt(0)
    r_f2 = f2.add_run()
    r_f2.font.name = 'Times New Roman'
    r_f2.font.size = Pt(10)
    fld_xml2 = '<w:fldSimple ' + nsdecls('w') + ' w:instr="PAGE"/>'
    r_f2._r.append(parse_xml(fld_xml2))
    
    pg_num_xml2 = '<w:pgNumType ' + nsdecls('w') + ' w:fmt="decimal" w:start="1"/>'
    s2._sectPr.append(parse_xml(pg_num_xml2))

    # ====================================================
    # CHAPTER ONE: BACKGROUND OF WEGAGEN BANK
    # ====================================================
    add_p("Chapter One: Background of Wegagen Bank", style='Heading 1', space_before=0, space_after=8)
    
    add_p("1.1 Introduction", style='Heading 2')
    add_p("Wegagen Bank Share Company is one of Ethiopia's leading private commercial banks, providing a wide range of financial services to individuals, businesses, and institutions across the country. In today's world, the banking sector is no longer just about deposits and loans — it is increasingly driven by technology. Digital banking, mobile applications, card payment systems, and automated core banking platforms have become essential to how banks operate and compete. While this digital transformation brings convenience and efficiency, it also introduces challenges such as system reliability, cybersecurity threats, and the constant need for software development and maintenance to keep services running smoothly.")
    add_p("Because of these demands, banks like Wegagen depend heavily on skilled software developers and IT professionals to build, maintain, and improve the systems that power everyday banking services — from mobile and internet banking to card payments and internal enterprise systems. Wegagen Bank's continuous investment in technology reflects its recognition that strong, secure, and innovative digital infrastructure is central to remaining competitive and serving millions of customers reliably.")

    add_p("1.2 Brief History of Wegagen Bank", style='Heading 2')
    add_p("Wegagen Bank S.C. was established on June 11, 1997, by sixteen visionary founders who recognized the important role financial institutions play in driving sustainable economic development. The Bank began operations with an initial paid-up capital of ETB 30 million and started with two branches — Goffa and Meskel — in Addis Ababa.")
    add_p('Over the years, the Bank expanded steadily, introducing new technologies and services to keep pace with a changing financial landscape. In August 2000, it became a pioneer in connecting its branches through a Wide Area Network (WAN). In 2010, it introduced card banking through its "Agar" Visa Debit Card, followed by mobile banking in 2014 and internet banking in 2018. In September 2017, the Bank inaugurated its 23-storey headquarters building on Ras Mekonnen Street, opposite Addis Ababa Stadium — the same building where my internship took place.')
    add_p('More recently, the Bank has continued to modernize its services, launching a digital lending platform called "Efoyta" in July 2024 and a Visa Prepaid International Card in August 2024. Today, Wegagen Bank is among Ethiopia\'s largest private financial institutions, operating 443 branches across the country under ten district offices, and serving more than 3 million customers.')

    p_hq_img = doc.add_paragraph()
    p_hq_img.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_hq_img.paragraph_format.space_before = Pt(6)
    p_hq_img.paragraph_format.space_after = Pt(2)
    p_hq_img.add_run().add_picture(io.BytesIO(imgs['rId9']), width=Inches(3.8))
    add_caption("Figure 1.1: Wegagen Bank Headquarters")

    add_p("1.3 Main Products and Services of Wegagen Bank", style='Heading 2')
    add_p("Wegagen Bank S.C. is one of Ethiopia’s leading private commercial banks, committed to driving financial inclusion, economic growth, and digital banking innovation across the country. Over the years, it has transformed into a customer-centric financial institution delivering a broad portfolio of conventional, interest-free, and technology-driven banking solutions. Some of its main products and services include:")

    add_p("1.3.1 Deposit and Savings Services", style='Heading 3')
    add_p("One of Wegagen Bank’s core offerings is providing safe and rewarding deposit avenues for individual and corporate clients. These include regular savings accounts, current (checking) accounts, fixed-time deposits, and tailored packages such as youth, women, and diaspora savings accounts. These services encourage a strong savings culture while giving customers flexible access to their funds.")

    add_p("1.3.2 Credit and Loan Facilities", style='Heading 3')
    add_p("The Bank supports businesses, agricultural enterprises, and individuals by extending various credit facilities. Its financing solutions include working capital loans, term loans, merchandise loans, import/export financing, trade credit, and personal consumer loans. Through responsible lending, the Bank plays a key role in stimulating commerce and industrial activity.")

    add_p("1.3.3 Digital and Electronic Banking (e-Banking)", style='Heading 3')
    add_p("To ensure 24/7 banking convenience, Wegagen Bank provides an extensive digital ecosystem:")
    add_bullet("Customers can check balances, transfer funds, pay utility bills, and manage accounts directly from their smartphones or web browsers.", bold_prefix="Mobile & Internet Banking: ")
    add_bullet("The Bank issues local proprietary debit cards and internationally co-branded Visa debit and prepaid cards, supported by a widespread network of ATMs and Point-of-Sale (POS) terminals across the country.", bold_prefix="Card Banking (Agar Visa): ")
    add_bullet("An innovative digital credit service that provides fast, collateral-free micro-loans and automated credit scoring for small businesses, gig workers, and individuals through mobile channels.", bold_prefix="Efoyta Digital Micro-Lending: ")

    add_p("1.3.4 Interest-Free Banking (Wegagen Amanah)", style='Heading 3')
    add_p("Recognizing the diverse financial needs of Ethiopian society, Wegagen Bank provides Sharia-compliant banking services under the brand name Amanah. Governed by strict Islamic jurisprudence, these include interest-free deposit accounts (Wadiah, Mudarabah) and financing mechanisms (Murabaha, Ijarah), offered through dedicated service windows as well as full-fledged independent Amanah branches.")

    add_p("1.4 Main Customers and End Users of Wegagen Bank's Products and Services", style='Heading 2')
    add_p("Wegagen Bank’s customer base spans from individual retail consumers to high-profile corporate conglomerates, institutional entities, and previously unbanked populations. The primary customer groups and their engagement with the Bank's financial solutions are broken down below:")

    add_p("1.4.1 Retail and Individual Banking Customers", style='Heading 3')
    add_p("Individual consumers represent the largest segment by volume, engaging daily with conventional and digital retail banking solutions.")
    add_bullet("Secure avenues for personal savings and wealth accumulation; Frictionless 24/7 access to account balances, utility bill payments, and P2P funds transfers.", bold_prefix="Key Needs: ")
    add_bullet("Rely on direct salary deposit accounts, personal consumer loans, and Agar Visa debit cards for Point-of-Sale (POS) and ATM transactions.", bold_prefix="Salaried Employees & Professionals: ")
    add_bullet("Utilize mobile banking channels and specialized youth savings accounts tailored to support early financial independence.", bold_prefix="Youth & University Students: ")
    add_p("Without accessible retail services, daily commerce and individual household financial management would face significant friction and cash dependency.")

    add_p("1.4.2 Micro, Small, and Medium Enterprises (MSMEs)", style='Heading 3')
    add_p("MSMEs form the backbone of the domestic economy, requiring flexible credit structures and working capital financing from Wegagen Bank.")
    add_bullet("Accessible short-term credit facilities to manage operational cash flow; Seamless merchant collection channels to receive customer payments.", bold_prefix="Key Needs: ")
    add_bullet("Working capital loans and overdraft facilities to finance inventory or raw materials; Merchant POS terminals and merchant QR code/mobile payment solutions for retail storefronts.", bold_prefix="Services Utilized: ")

    add_p("1.4.3 Corporate and Commercial Entities", style='Heading 3')
    add_p("Large private corporations, commercial manufacturers, transport companies, and import/export enterprises depend on Wegagen Bank’s institutional capital base and trade finance infrastructure.")
    add_bullet("High-volume capital lending for long-term industrial projects and expansion; Facilitation of international trade and cross-border currency clearance.", bold_prefix="Key Needs: ")
    add_bullet("Secure multi-user approval workflows for bulk payroll, vendor disbursements, and real-time treasury monitoring.", bold_prefix="Corporate Internet Banking: ")
    add_bullet("Acquisition of foreign currency for raw material importation and overseas supply chain settlement.", bold_prefix="Foreign Exchange Services: ")

    add_p("1.5 Organization and Structure", style='Heading 2')
    add_p("Wegagen Bank operates under a Board of Directors and a Core Management team led by the Chief Executive Officer, supported by three Deputy Chief Executive Officers overseeing key functional areas: Operations, Enterprise, and Technology. Under the Technology function, the Bank maintains several specialized units responsible for keeping its digital and core banking systems running securely and efficiently. Some of these units include:")
    add_bullet("Responsible for managing, maintaining, customizing, and supporting the Bank's core banking system (Oracle Flexcube) and other critical enterprise applications that power day-to-day banking operations. (This is where I worked during my internship.)", bold_prefix="Core Application Management Services – ")
    add_bullet("Manages the Bank's networks, servers, and cybersecurity systems.", bold_prefix="IT Infrastructure and Security Management – ")
    add_bullet("Oversees mobile, internet, agent, and card banking platforms.", bold_prefix="Digital Banking Operations – ")
    add_bullet("Support the Bank's overall operations and workforce.", bold_prefix="Human Resource Management, Finance, and other administrative departments – ")
    add_p("The Core Application Management Services department, where I was placed, plays a central role in ensuring that the Bank's core systems and supporting applications function reliably, are properly maintained, and are continuously improved to meet the Bank's growing operational and customer service needs.")

    add_p("1.6 Workflow in Core Application Management Services", style='Heading 2')
    add_p("The workflow within the Core Application Management Services department follows a clear and practical process designed to resolve issues quickly while maintaining the stability of the Bank's core systems. For example:")
    add_bullet("An issue, bug, or system-related problem is reported to the department, typically through a phone call to the office from staff experiencing the problem.")
    add_bullet("The team receives the report and takes note of the issue, prioritizing it based on urgency and impact on banking operations.")
    add_bullet("The reported issue is discussed during a team meeting, where it is evaluated and assigned to the appropriate team member(s) for resolution.")
    add_bullet("The assigned team member investigates the issue and develops or applies a fix.")
    add_bullet("Before any change is deployed to the live (production) system, it is first tested in a dedicated test environment to confirm that it resolves the issue without introducing new problems.")
    add_bullet("Once the fix passes testing, it is approved and deployed to the live system, restoring normal functionality.")
    add_bullet("If necessary, the department documents the issue and resolution for future reference, helping the team respond more efficiently to similar problems.")

    # ====================================================
    # CHAPTER TWO: PRACTICAL ATTACHMENT EXPERIENCE
    # ====================================================
    doc.add_page_break()
    add_p("Chapter Two: My Practical Attachment Experience", style='Heading 1', space_before=0, space_after=8)

    add_p("2.1 How I Joined the Company", style='Heading 2')
    add_p("I joined Wegagen Bank through the university's Practical Attachment program, which is designed to connect students with organizations relevant to their field of study. After submitting the application letter from my university, Wegagen Bank responded and I was invited to report to their headquarters, located on Ras Mekonnen Street, Addis Ababa.")
    add_p("When I arrived, I was introduced to the Bank's structure and given an overview of its different departments and how they support the Bank's operations. Given my background in software development, I was placed within the Core Application Management Services department, under the Bank's Technology function — a unit responsible for managing, maintaining, and developing the core banking system and other critical enterprise applications used across the Bank.")

    add_p("2.1.1 Placement within Wegagen Bank: Core Application Management Service", style='Heading 3')
    add_p("From June 23, 2026, to September 1, 2026, I was placed within the Core Application Management Service under the Information Technology Directorate of Wegagen Bank S.C., stationed at the Bank’s headquarters in Addis Ababa. This department serves as one of the most critical technological pillars of the Bank, tasked with maintaining, optimizing, and extending core digital banking architectures, enterprise service integrations, and customer-facing digital channels that power the institution's day-to-day operations.")
    add_p("Within this department, my assignment centered on Mobile and Web Banking Application Development, specifically building modern, high-performance interfaces and scalable backend services to elevate the Bank's digital service delivery. The digital banking stream functions as a frontline engine for financial inclusion and customer retention by:")
    add_bullet('Developing a mobile application called "Smart Agent and Branch Queue Management." This project was introduced to me as a solution to a real, everyday problem at the Bank: long customer queues at branches, which waste customers\' time and create inefficiencies for branch staff.', bold_prefix="Engineering Responsive Digital Touchpoints: ")
    add_bullet("Developing secure, scalable RESTful application programming interfaces (APIs) capable of facilitating seamless business logic and secure data exchange.", bold_prefix="Building Resilient Backend Services: ")
    add_p("This functional area is particularly strategic within Wegagen Bank because it bridges enterprise core backend databases with direct consumer touchpoints, ensuring that account holders experience reliable, secure, and responsive digital interactions.")

    add_p("2.1.2 Supervision and Technical Leadership", style='Heading 3')
    add_p("My internship engagement was carried out under the close mentorship and direct technical guidance of Mr. Habtamu Assegahegn, Senior Software Engineer, who served as my primary supervisor throughout the deployment period.")
    add_p("Mr. Habtamu provided critical direction in translating enterprise banking workflows into modular, scalable technical requirements. Under his mentorship, I adhered to rigorous software engineering standards, established secure architectural workflows, and integrated real-world fintech design patterns.")

    add_p("2.2 Workflow in the Core Application Management Service Department", style='Heading 2')
    add_p("The Core Application Management Service Department consists of professionals working on different software development, system management, and application-support activities. Team members are assigned to different projects and tasks based on project requirements and organizational priorities.")
    add_p("During my internship, I worked on a mobile application project under the guidance and supervision of my assigned supervisor. The supervisor provided direction, assigned tasks, reviewed my progress, and gave feedback when necessary.")

    add_p("2.3 Work Tasks Executed During the Practical Attachment", style='Heading 2')
    add_p("During my practical attachment at Wegagen Bank, I worked on different software development activities under the guidance of my assigned supervisor. The internship provided me with an opportunity to apply my academic knowledge to a practical project that addresses a real-world problem in the banking sector.")
    add_p("My major project was the development of a Smart Banking Queue and Forex Finder System. The main purpose of the system was to reduce the inconvenience caused by long queues at bank branches by allowing customers to check queue information and hold their position remotely before visiting a branch. In addition, the platform was designed to provide customers with foreign-exchange rate information, allowing them to view available currency rates conveniently.")
    add_p("Before beginning implementation, I analyzed the problem, identified the expected users and system requirements, prepared the system design, and selected suitable technologies. Throughout the internship, I worked collaboratively under the supervision of my supervisor, received feedback, solved technical challenges, and documented my progress. The major tasks performed during the internship are described below.")

    add_p("2.3.1 Problem Identification and Requirement Analysis", style='Heading 3')
    add_p("At the beginning of the project, I focused on understanding the problems faced by customers when visiting bank branches. One of the major problems identified was the long waiting time experienced by customers who need banking services such as cash deposits, foreign-exchange services, and other branch-related services.")
    add_p("To address this problem, I proposed a system that would allow customers to view branch information, check the current queue, obtain an estimated waiting time, and join or hold a queue remotely. The platform also included a foreign-exchange feature that would allow customers to view currency rates such as the US dollar, euro, and British pound against the Ethiopian birr. The main users identified for the system included customers, bank employees or agents, and administrators. This stage helped me understand the project scope and establish a clear direction for development.")

    add_p("2.3.2 System Planning and Design", style='Heading 3')
    add_p("After identifying the requirements, I worked on the initial planning and design of the Smart Banking Queue and Forex Finder System. The system was planned as a mobile-based platform that communicates with a backend server through APIs. The mobile application would provide the user interface, while the backend would manage the business logic, user information, queue data, branch information, and exchange-rate data.")
    add_p("During this stage, I designed the overall system architecture and planned how the different components would communicate with one another. I also designed the customer queue workflow, including branch selection, service selection, queue joining, digital-token generation, and queue-status monitoring. In addition, I planned the foreign-exchange rate interface and considered the different access permissions required for customers, employees, and administrators. The proposed design was reviewed with my supervisor before I proceeded with implementation.")

    add_p("2.3.3 Backend Architecture and Database Implementation", style='Heading 3')
    add_p("The server-side infrastructure and data layer of the Smart Banking Queue and Forex Finder System were architected and implemented using Node.js and Express.js in conjunction with PostgreSQL as the core relational database management system, mediated by Prisma ORM. The backend serves as the centralized orchestration engine responsible for processing incoming client requests from the mobile interface, executing business logic, enforcing strict relational integrity, and serving structured, type-safe RESTful API endpoints. To ensure maintainability, scalability, and clean separation of concerns, the backend was organized around a layered architectural pattern separating the routing mechanisms, controllers, business services, and database persistence layers.")
    add_p("User identity verification and access management were integrated via Clerk as a third-party authentication provider, where incoming HTTP requests are intercepted by custom Express middleware to cryptographically validate session bearer tokens prior to reaching protected routes. Database modeling commenced by defining core business entities, their scalar attributes, and their relational mappings within the Prisma schema—encompassing users, branch locations, operational banking services, active queue entries, digital tokens, foreign exchange rates, notifications, and historical queue telemetry. PostgreSQL was utilized to enforce referential integrity and consistency through primary keys, foreign key constraints, unique indexing, and non-null column definitions, while Prisma Client facilitated type-safe query generation and automated schema migrations. When a user joins the queue, the system generates a unique sequential digital token, logs the active queue state, and persists the transaction to queue history. Comprehensive request validation schemas, centralized exception-handling middleware, and standardized HTTP response codes (such as 200 OK, 201 Created, 400 Bad Request, 401 Unauthorized, and 500 Internal Server Error) were incorporated across all foreign exchange, branch discovery, and queue endpoints, yielding a robust, high-integrity server-side system capable of reliable data exchange with the client application.")

    add_p("2.3.4 Mobile Application Development", style='Heading 3')
    add_p("The client-facing mobile application was architected and implemented using React Native and TypeScript powered by the Expo framework, functioning as the primary user touchpoint for branch discovery, remote queue reservations, real-time ticket tracking, and foreign exchange telemetry. Development commenced with initializing the Expo environment and establishing a file-based routing architecture via Expo Router, creating an intuitive screen hierarchy encompassing authentication flows, dashboard landing views, branch and agent locators, service selection modules, live queue status monitors, digital token displays, profile configurations, and historical queue records.")
    add_p("Identity management, biometrics, and session persistence were integrated via the Clerk Expo SDK (@clerk/clerk-expo) coupled with Expo SecureStore and device-level biometric authentication (such as fingerprint recognition and facial scanning), enabling secure hardware-backed credential storage and frictionless biometric re-authentication alongside standard login protocols. The presentation layer was styled systematically using React Native's native StyleSheet.create API, standardizing responsive layouts, component hierarchy, typographic scales, spacing systems, and platform-specific visual ergonomics across varying viewport dimensions. For the client-server data synchronization layer, an Axios HTTP client was configured with request interceptors to automatically inject bearer authorization tokens retrieved from SecureStore into request headers, enabling secure data transport to protected backend REST endpoints. The system also supports real-time notifications using Socket.IO, in which customers and agents are notified instantly when an event occurs.")

    # Figures 2.1 (Two side-by-side screenshots)
    p_app_imgs = doc.add_paragraph()
    p_app_imgs.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_app_imgs.paragraph_format.space_before = Pt(6)
    p_app_imgs.paragraph_format.space_after = Pt(2)
    p_app_imgs.add_run().add_picture(io.BytesIO(imgs['rId10']), width=Inches(2.5))
    p_app_imgs.add_run("   ")
    p_app_imgs.add_run().add_picture(io.BytesIO(imgs['rId11']), width=Inches(2.5))
    add_caption("Figure 2.1: Smart Agent Mobile App – Home and Branch Discovery Screens")

    add_p("To support in-branch operations, dedicated teller and employee interface screens were designed to allow bank personnel to actively manage customer throughput, view assigned branch service queues, monitor active ticket counters, and advance the queue line by invoking API actions to serve and call the next waiting customer in real time.")

    # Figure 2.2 (Teller dashboard)
    p_teller_img = doc.add_paragraph()
    p_teller_img.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_teller_img.paragraph_format.space_before = Pt(6)
    p_teller_img.paragraph_format.space_after = Pt(2)
    p_teller_img.add_run().add_picture(io.BytesIO(imgs['rId12']), width=Inches(2.8))
    add_caption("Figure 2.2: Dedicated Teller and Employee Queue Management Interface")

    add_p("2.3.5 Queue Management and Foreign-Exchange Rate Implementation", style='Heading 3')
    add_p("The core transactional capabilities of the system were established through the integration of an automated remote queue-management pipeline and a real-time foreign-exchange (forex) rate finder, engineered to streamline in-branch service throughput and provide transparency for currency operations. The queue-management module mitigates physical branch congestion by enabling users to locate a target branch, select a designated banking service, inspect real-time queue length metrics, and receive dynamic wait-time estimations prior to physical arrival.")
    add_p("Upon initiating a remote reservation, the backend atomically assigns an indexed digital token representing the user's explicit position, broadcasting status updates to the client interface so users can track queue progression while recording completed, cancelled, or expired ticket states into the persistent queue history for operational analytics and auditability. Concurrently, the foreign-exchange rate finder delivers up-to-date currency valuations against the Ethiopian Birr (ETB)—specifically targeting high-volume foreign currencies such as the US Dollar (USD), Euro (EUR), and British Pound (GBP)—by synchronizing with backend rate services to render structured rate matrices detailing distinct transaction types, including buying and selling margins for both physical cash notes and bank-transfer instruments. The administrator can update the foreign currency exchange rates based on the bank's daily rate-setting policies.")

    add_p("2.3.6 Quality Assurance, System Documentation, and Project Delivery", style='Heading 3')
    add_p("API standardization, comprehensive testing, multi-stage debugging, and professional documentation workflows were executed systematically throughout the development lifecycle to ensure platform reliability, cross-stack interoperability, and structured project governance. To establish a dependable integration contract between the mobile client and backend services, RESTful endpoints were defined and cataloged using Swagger/OpenAPI specifications, providing detailed schemas for HTTP request methods, path and query parameters, JSON payload structures, Clerk bearer token authentication requirements, validation rules, and standardized HTTP response codes.")
    add_p("Systematic endpoint testing, database transactional verification, and integration stress tests were performed across core functional modules—including user authentication lifecycles, atomic queue token issuance, real-time wait telemetry, and foreign-exchange data retrieval—allowing me to identify and resolve critical edge cases involving network latency, environment configuration mismatches, asynchronous data races, and payload serialization errors. In parallel with technical testing, the development process was accompanied by rigorous documentation and engineering reporting.")

    # Figure 2.3 (Swagger UI screenshots)
    p_swg_img = doc.add_paragraph()
    p_swg_img.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_swg_img.paragraph_format.space_before = Pt(6)
    p_swg_img.paragraph_format.space_after = Pt(2)
    p_swg_img.add_run().add_picture(io.BytesIO(imgs['rId13']), width=Inches(5.0))
    p_swg2_img = doc.add_paragraph()
    p_swg2_img.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_swg2_img.paragraph_format.space_before = Pt(2)
    p_swg2_img.paragraph_format.space_after = Pt(2)
    p_swg2_img.add_run().add_picture(io.BytesIO(imgs['rId14']), width=Inches(5.0))
    add_caption("Figure 2.3: Backend REST API Swagger/OpenAPI Documentation")

    add_p("2.4 Materials and Methodology", style='Heading 2')
    add_p("During my practical attachment at Wegagen Bank, I used different hardware, software tools, frameworks, libraries, and development technologies to design and implement the Smart Banking Queue and Forex Finder System. The project was developed as a mobile-based application supported by a backend server, a PostgreSQL database, and an administrative interface. These tools helped me implement the system’s main functions, including user authentication, branch and service management, digital queue joining, token generation, real-time queue updates, notifications, foreign-exchange rate display, and administrative management.")

    add_p("2.4.1 Development Environment and Tools", style='Heading 3')
    add_p("The system was developed using a personal computer running Windows 11. Visual Studio Code was used as the primary integrated development environment for writing, organizing, and debugging the application’s source code. Expo Go was used to run and test the React Native mobile application on a physical mobile device during development.")
    add_p("PostgreSQL was used as the relational database management system for storing and managing the system’s data. pgAdmin was used to inspect and manage the PostgreSQL database, execute database-related operations, and verify stored records.")

    add_p("2.4.2 Frontend and Mobile Technologies", style='Heading 3')
    add_bullet("React Native was the main framework used to develop the mobile application. It enabled me to build the Smart Banking Queue and Forex Finder System for mobile devices using TypeScript-based technologies while maintaining a component-based development approach. Reusable components and clean screen separation were maintained throughout.", bold_prefix="React Native: ")
    add_bullet("Expo simplified the development and testing process by providing robust tooling for running the application and accessing native device features. Expo Go was particularly useful for real-time testing on physical Android devices.", bold_prefix="Expo & Expo Go: ")
    add_bullet("Expo Router was used to manage navigation and routing within the React Native application, utilizing a clean, file-based routing structure.", bold_prefix="Expo Router: ")
    add_bullet("TypeScript was used across both the frontend and backend codebases, providing static type safety and catching potential runtime errors during development.", bold_prefix="TypeScript: ")

    add_p("2.4.3 Backend and Database Technologies", style='Heading 3')
    add_bullet("Node.js was used as the runtime environment for executing server-side TypeScript code and handling client requests asynchronously.", bold_prefix="Node.js: ")
    add_bullet("Express.js was used as the server web framework to implement RESTful APIs, route requests, execute business middleware, and format JSON responses.", bold_prefix="Express.js: ")
    add_bullet("PostgreSQL served as the relational database, enforcing data integrity, foreign keys, unique constraints, and ACID-compliant transactional consistency.", bold_prefix="PostgreSQL: ")
    add_bullet("Prisma ORM connected the Node.js backend to the PostgreSQL database, providing type-safe query generation, automated schema migrations, and intuitive entity modeling.", bold_prefix="Prisma ORM: ")

    add_p("2.4.4 Authentication, Real-Time Communication, and Testing", style='Heading 3')
    add_bullet("Clerk provided user authentication, identity verification, session persistence, and email notification mechanisms for both customers and bank staff.", bold_prefix="Clerk Authentication: ")
    add_bullet("Socket.IO enabled bidirectional real-time communication between the server and connected mobile clients, broadcasting live queue status changes instantly.", bold_prefix="Socket.IO: ")
    add_bullet("Bruno was used as the primary API testing client to send requests, inspect responses, test error conditions, and validate backend endpoints before frontend integration.", bold_prefix="Bruno API Client: ")

    add_p("2.4.5 Version Control and Methodology", style='Heading 3')
    add_bullet("Git was used for tracking source code changes, creating commits, managing branches, and maintaining version history.", bold_prefix="Git: ")
    add_bullet("GitHub hosted the remote repository, providing code backup, structured issue tracking, and project organization.", bold_prefix="GitHub: ")
    add_p("Development followed an iterative, sequential methodology: requirement analysis $\\rightarrow$ system architecture and database design $\\rightarrow$ backend API development and validation $\\rightarrow$ frontend mobile implementation $\\rightarrow$ real-time integration $\\rightarrow$ comprehensive end-to-end testing with Bruno, Expo Go, and pgAdmin.")

    add_p("2.5 Evaluation of Work Performance", style='Heading 2')
    add_p("At the beginning of my practical attachment, several aspects of the development process were new to me. Although I had some previous experience with React, React Native was new, and I needed time to understand mobile application development and become familiar with the tools and technologies used in the project. Working with technologies such as Expo, Prisma, Clerk, and Socket.IO also required additional learning and practice. I mainly relied on official documentation to understand these technologies and apply them correctly during development.")
    add_p("In addition to the technical challenges, adjusting to the internship routine required commitment and proper time management. Since the internship workplace was in Addis Ababa and I live in Debre Zeit/Bishoftu, I had to wake up at around 5:00 a.m. to prepare and arrive at the office on time. Maintaining this routine while learning new technologies and completing development tasks helped me improve my discipline, punctuality, and sense of responsibility.")
    add_p("As the project was developed individually, managing the entire development process was one of the main responsibilities I handled. I worked on different parts of the system, including the backend APIs, database models, authentication, mobile application, real-time communication, queue management, foreign-exchange rate functionality, and administrative interface. This required me to organize the work carefully, understand how the different components were connected, and solve problems independently.")
    add_p("I followed a daily and weekly work plan to organize my tasks and maintain progress. I divided the project into smaller activities and worked on the different features step by step. I also used Bruno to test the backend APIs and verify whether the endpoints were working as expected before and after integrating them with the frontend.")
    add_p("Throughout the internship, I received positive feedback from my supervisor, who was pleased with the progress and functionality of my work. He also suggested improvements that helped me refine the system and improve certain aspects of its design and implementation. I considered this feedback during the development process and made adjustments where necessary.")
    add_p("By the end of the practical attachment, I had completed the planned project features and gained practical experience in developing a full-stack application. I was able to design relational database models, develop RESTful APIs, integrate authentication, implement real-time communication, and build both customer and administrative interfaces. Overall, the internship improved my technical knowledge, problem-solving ability, independence, time management, and confidence in developing software systems.")

    add_p("2.6 Challenges Faced During the Practical Attachment", style='Heading 2')
    add_p('During my time at Wegagen Bank, I faced several challenges that tested my ability to learn, adapt, and problem-solve. Unlike working within an established codebase, I was building the "Smart Agent and Branch Queue Management" application largely from scratch, which brought its own unique set of difficulties — from learning an unfamiliar framework to making architectural decisions with no existing code to guide me. Each challenge, however, taught me something valuable and contributed to my growth as a developer throughout the internship.')

    add_p("2.6.1 Learning React Native and Expo from the Ground Up", style='Heading 3')
    add_p("The most significant and persistent challenge I faced was learning React Native and Expo, a mobile development framework I had little prior hands-on experience with. While I had some general programming background with React, building a real mobile application required an entirely different mindset from web development — understanding how components render on mobile devices, how navigation works between screens, how to handle mobile-specific behavior, and how Expo's tooling fit into the development and testing workflow.")
    add_p("Because there was no existing codebase to learn from, I could not simply read through someone else's implementation to understand how things were supposed to work. Instead, I had to learn the framework's concepts directly — through documentation, experimentation, and trial and error — while simultaneously deciding how to structure the application myself.")

    add_p("2.6.2 Designing the Application Architecture from Scratch", style='Heading 3')
    add_p("Since the application was built from the ground up, I was also responsible for making foundational decisions that a developer joining an existing project would not normally have to make — how to structure the frontend screens, how the backend APIs should be organized, and how data would flow between the customer-facing app, the queue system, and agent assignment logic. Without an established pattern to follow, I had to think carefully about these decisions, knowing that mistakes made early in the architecture could cause bigger problems later in development.")

    add_p("2.6.3 API Integration and Client-Server Communication Bug", style='Heading 3')
    add_p("The most notable technical challenge I faced was connecting the React Native frontend to the Express/Node.js backend using Axios. Even though both sides of the application were built by me, getting them to communicate correctly was far from straightforward. Requests sent from the mobile app were not always reaching the backend as expected, and responses were not always being handled correctly on the frontend side — issues that could stem from anything as small as an incorrect endpoint URL, mismatched data formats between the request and what the backend expected, or how the app handled asynchronous responses.")
    add_p("Resolving this required systematically checking each part of the communication chain — verifying the backend endpoints were working correctly on their own, confirming the frontend was sending requests in the right format, and tracing how responses were being received and processed in the app. When I could not resolve the issue on my own, I brought it to my supervisor, Mr. Habtamu Assegahegn, who helped me pinpoint the exact cause and understand the correct way to structure the communication between the two sides of the application.")

    add_p("2.6.4 Balancing Learning with a Demanding Commute and Schedule", style='Heading 3')
    add_p("Alongside these technical challenges, I also had to manage the practical demands of the internship itself. Waking up around 5:00 AM each day to reach the office on time because the office was located in Addis Ababa, and I live in Debre Zeit, while simultaneously learning a new framework and building an application from scratch, meant I had to be disciplined about how I used my time and energy each day to stay productive and keep making progress.")

    add_p("2.7 Measures Taken to Overcome Challenges", style='Heading 2')
    add_p("Despite the challenges I encountered during my practical attachment, I was able to work through each of them by combining patience, self-study, systematic investigation, and guidance from my mentor.")

    add_p("2.7.1 Overcoming the React Native and Expo Learning Curve", style='Heading 3')
    add_p("I approached this by being patient and systematic rather than trying to learn everything about the framework at once. I focused on understanding one concept at a time — navigation, screen components, state handling — starting with whatever I needed for the specific feature I was currently building. I relied on the official React Native and Expo documentation as my primary reference, reading through core concepts and then immediately applying them in the project so they became concrete rather than abstract.")

    add_p("2.7.2 Iterative Design and Incremental Architecture", style='Heading 3')
    add_p("Since there was no existing codebase to guide me, I dealt with this by starting simple and refining as I went, rather than trying to design a perfect structure from the beginning. I built the core functionality first — joining a queue, viewing status, assigning agents — and then improved the structure of both the frontend and backend as my understanding of the requirements became clearer.")

    add_p("2.7.3 Systematic Debugging of API Endpoints", style='Heading 3')
    add_p("To resolve the issue connecting the React Native frontend to the Express/Node.js backend, I traced through the communication chain step by step rather than guessing at a quick fix. I checked the backend endpoints independently using Bruno to confirm they worked correctly on their own, then verified how the frontend was formatting and sending its requests, and finally checked how responses were being received and handled in the app. This systematic process helped me isolate exactly where the breakdown was happening. When I reached a point I could not resolve alone, I brought the specific issue to my supervisor, Mr. Habtamu, who helped me identify the exact cause and guided me toward the correct fix.")

    add_p("2.7.4 Time Management and Structured Routine", style='Heading 3')
    add_p("To handle the early mornings and the pressure of learning while building at the same time, I stayed disciplined about my daily routine — preparing the night before and prioritizing my most demanding learning or coding tasks for when I had the most energy and focus. This consistency helped me stay productive throughout the internship despite the early start times.")

    # ====================================================
    # CHAPTER THREE: BENEFITS GAINED
    # ====================================================
    doc.add_page_break()
    add_p("Chapter Three: Benefits Gained", style='Heading 1', space_before=0, space_after=8)

    add_p("3.1 Practical Skills", style='Heading 2')
    add_p("My internship at Wegagen Bank significantly strengthened my practical, hands-on development skills, particularly in mobile application development and full-stack engineering.")
    add_p("One of the most important areas of growth was my proficiency in React Native and Expo. Coming into the internship with limited experience in mobile development, I had to learn the framework largely on my own while simultaneously building a real application. By the end of the internship, I was comfortable structuring navigation between screens, managing component state, handling asynchronous data, and building a functional, user-facing mobile interface from the ground up.")
    add_p("I also gained meaningful practical experience with TypeScript, learning how static typing helps catch errors earlier in development and makes code easier to maintain. On the backend side, I developed stronger skills in building server-side applications using Express.js and Node.js.")

    add_p("3.2 Theoretical Understanding Gained", style='Heading 2')
    add_p("Alongside the practical skills, my internship deepened my theoretical understanding of several important areas of software engineering and system design.")

    add_p("3.2.1 Mobile Application Architecture", style='Heading 3')
    add_p("Working with React Native and Expo taught me the underlying concepts of how mobile applications are structured differently from web applications — how components render on-device, how navigation and screen lifecycles work, and how mobile apps communicate with backend servers over the network. This gave me a clearer theoretical picture of the client-server model in a mobile context, not just in theory but through building it myself.")

    add_p("3.2.2 API Design and Client-Server Communication", style='Heading 3')
    add_p("Through building and debugging the backend APIs for the queue management system, I gained a much stronger theoretical understanding of how client-server communication works — request/response cycles, data formatting (such as JSON payloads), endpoint design, and the importance of consistent contracts between frontend and backend. The API integration challenge I faced was, in many ways, a practical lesson in this theory: I had to understand exactly how HTTP requests are constructed, sent, and interpreted to figure out where the communication was breaking down.")

    add_p("3.2.3 Software Development in a Real-World, Production-Oriented Environment", style='Heading 3')
    add_p("Unlike academic assignments, developing an application meant for real use inside a bank required me to think about reliability, correctness, and maintainability — not just whether the code worked once, but whether it would keep working under real conditions. This gave me a deeper appreciation for good software design principles, including separating frontend and backend responsibilities clearly and writing code that could be understood and extended later.")

    add_p("3.3 Interpersonal and Communication Skills Gained", style='Heading 2')
    add_p("One of the most valuable non-technical outcomes of my internship was the growth in my interpersonal and communication skills, developed mainly through my regular interaction with my supervisor, Mr. Habtamu Assegahegn, and other staff within the department.")
    add_p("At the beginning, explaining technical progress and problems to a supervisor with far more professional experience than me felt intimidating. I was expected to regularly present my progress on the queue management application — what I had built, what was working, and what challenges I was facing. In the early weeks, I sometimes struggled to explain technical issues clearly and concisely, especially when describing bugs like the API integration problem.")
    add_p("Over time, however, these regular presentations became a valuable exercise in communication. I learned how to summarize technical work in a way that was clear and easy to follow, how to describe a problem specifically rather than vaguely, and how to ask for help in a focused, useful way rather than simply saying 'it doesn\\'t work.' Receiving regular feedback from my supervisor also taught me how to listen actively, accept constructive criticism without becoming discouraged, and apply that feedback immediately to improve my next piece of work. By the end of the internship, presenting my progress and discussing technical challenges felt far more natural than it had at the start.")

    add_p("3.4 Teamwork Skills", style='Heading 2')
    add_p("Although much of my development work was done independently, my internship still gave me valuable experience in working as part of a larger team within the Core Application Management Services department. Being placed inside a professional department meant learning how individual work fits into a bigger, shared responsibility — the department's role of keeping the Bank's core systems running reliably.")
    add_p("I saw firsthand how the department handled incoming issues as a team: problems reported by phone were discussed together in team meetings, assigned based on availability and expertise, and tracked through to resolution. Being part of these meetings, even as an intern, taught me how a technical team coordinates around shared priorities and communicates about ongoing work.")

    add_p("3.5 Understanding of Work Ethics", style='Heading 2')
    add_p("My internship at Wegagen Bank also gave me a much clearer understanding of professional work ethics, shaped by the seriousness and structure of working inside a national financial institution.")
    add_bullet("Reaching the headquarters on time every day required me to wake up around 5:00 AM, since I needed to travel from Debre Zeit to Addis Ababa to arrive at the office punctually. This daily discipline taught me that punctuality is not just a personal habit but a form of respect for colleagues, meetings, and deadlines — arriving late or unprepared would have disrupted both my own progress and the team's plans.", bold_prefix="Punctuality and Discipline: ")
    add_bullet("Because I was building an application largely on my own, I quickly learned that the quality and progress of my work depended entirely on my own effort and initiative. Regularly presenting my progress to my supervisor meant I had to take ownership of both my successes and the problems I encountered, rather than hiding difficulties or making excuses.", bold_prefix="Accountability and Ownership: ")
    add_bullet("Working within a bank's technology department also made me conscious of the importance of confidentiality and responsible handling of systems and data. Since Core Application Management Services deals directly with core banking systems, I understood the need to be careful, professional, and trustworthy in how I approached my work.", bold_prefix="Confidentiality and Professionalism: ")

    # ====================================================
    # CHAPTER FOUR: CONCLUSION AND RECOMMENDATIONS
    # ====================================================
    doc.add_page_break()
    add_p("Chapter Four: Conclusion and Recommendations", style='Heading 1', space_before=0, space_after=8)

    add_p("4.1 Conclusion", style='Heading 2')
    add_p("My two-and-a-half-month internship at Wegagen Bank S.C., within the Core Application Management Services department, was a genuinely transformative experience that helped me connect the theoretical foundations of my Computer Science education with the practical realities of building software for a real financial institution. Starting from an initial phase of studying how quality software is delivered and how the bank manages its core systems, and moving into the hands-on design and development of the Smart Agent & Branch Queue and Forex Finder mobile application, I came to understand that software is not built overnight — it requires a deep understanding of the problem being solved, the constraints of the environment it will run in, and the people who will ultimately use it. Building the backend before the frontend, structuring routes, middleware, and services carefully, and only then layering the user interface on top, taught me that solid software is the product of deliberate sequencing and patience rather than speed.")
    add_p("Working inside a bank, rather than a pure software company, also gave me an appreciation for how policy and process shape technical decisions — understanding, for instance, why the forex-rate feature had to follow the bank's internal daily-update policy rather than pulling from a live external feed, and why every change to a customer-facing system has to be considered carefully in a regulated, trust-dependent environment. This showed me that technical skill alone is not enough; understanding the organization and the rules it operates under is just as important to delivering software that actually works in context.")
    add_p("Technically, the internship deepened my abilities across the stack — from Node.js, Express, and PostgreSQL on the backend, to React Native and TypeScript on the frontend. But more than any single technology, it was the discipline of building the backend from the ground up, routing, database connections, middleware, and services, that I found myself drawn to, and it left me with a genuine interest in pursuing backend development further.")

    add_p("4.2 Recommendations", style='Heading 2')
    add_p("Based on my internship experience, I would like to offer the following recommendations to improve the program for future students and strengthen the relationship between the university and the industry.")

    add_p("4.2.1 Structured Onboarding for Interns", style='Heading 3')
    add_p("One of the early challenges I faced was that the first weeks were spent largely on self-directed study — learning how quality software is delivered and how the bank manages its core systems — without a structured, written introduction to the department's tools, coding standards, or expectations. Host departments could shorten this ramp-up period by providing a brief onboarding document covering the development environment setup, the tools and frameworks in use, and the standards a new intern's code is expected to follow. Even a short written guide would let interns move into meaningful, hands-on work sooner.")

    add_p("4.2.2 Pre-Internship Skill Preparation", style='Heading 3')
    add_p("Several of the technologies I used during the internship — Node.js with Express on the backend, PostgreSQL, React Native with Expo, and third-party services like Clerk for authentication — were tools I had to learn largely on the job. If the university offered short, practical sessions on backend frameworks, mobile app development, database integration, and version control before students begin their internships, students would arrive better prepared and be able to contribute confidently from the very first weeks rather than spending a large portion of the internship on foundational learning.")

    add_p("4.2.3 University and Industry Collaboration", style='Heading 3')
    add_p("There is currently a noticeable gap between what is taught in the classroom and what is expected in a real, policy-governed environment such as a bank's software department. Strengthening ties between the university and organizations like Wegagen Bank — through guest lectures, joint workshops, or regular communication about which skills are in demand — would help keep the curriculum relevant and better prepare students for the realities of professional, enterprise-grade software development.")

    # ====================================================
    # REFERENCES
    # ====================================================
    doc.add_page_break()
    add_p("References", style='Heading 1', space_before=0, space_after=8)
    
    refs = [
        "Wegagen Bank S.C. (2025). Prospectus for registration of shares currently held by shareholders. Ethiopian Securities Exchange. https://esx.et/wp-content/uploads/2025/03/Wegagen_Bank_S_C_Prospectus-for-Registration-of-Shares-Currently-Held-by-Shareholders.pdf",
        "Wegagen Bank S.C. (2026). Official website. https://wegagen.com/",
        "Wikipedia contributors. (2026). Wegagen Bank. Wikipedia. https://en.wikipedia.org/wiki/Wegagen_Bank",
        "Meta Platforms, Inc. (2026). React Native documentation. https://reactnative.dev/docs/getting-started",
        "Expo. (2026). Expo documentation. https://docs.expo.dev/",
        "OpenJS Foundation. (2026). Node.js documentation. https://nodejs.org/en/docs",
        "OpenJS Foundation. (2026). Express.js — Node.js web application framework. https://expressjs.com/",
        "PostgreSQL Global Development Group. (2026). PostgreSQL documentation. https://www.postgresql.org/docs/",
        "Clerk, Inc. (2026). Clerk documentation. https://clerk.com/docs"
    ]
    
    for r in refs:
        p_ref = doc.add_paragraph()
        p_ref.paragraph_format.left_indent = Inches(0.5)
        p_ref.paragraph_format.first_line_indent = Inches(-0.5)
        p_ref.paragraph_format.space_before = Pt(0)
        p_ref.paragraph_format.space_after = Pt(4)
        p_ref.paragraph_format.line_spacing = 1.15
        p_ref.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        r_run = p_ref.add_run(r)
        r_run.font.name = 'Times New Roman'
        r_run.font.size = Pt(11)

    # Save document
    dst_path = "Wegagen Bank Sc(software development).docx"
    doc.save(dst_path)
    print("Document successfully generated and saved to", dst_path)

    # Enable updateFields in settings.xml
    temp_dir = tempfile.mkdtemp()
    with zipfile.ZipFile(dst_path, 'r') as zin:
        zin.extractall(temp_dir)
    
    settings_xml_path = os.path.join(temp_dir, 'word', 'settings.xml')
    if os.path.exists(settings_xml_path):
        with open(settings_xml_path, 'r', encoding='utf-8') as f:
            content = f.read()
        if '<w:updateFields' not in content:
            content = content.replace('</w:settings>', '<w:updateFields w:val="true"/></w:settings>')
            with open(settings_xml_path, 'w', encoding='utf-8') as f:
                f.write(content)
            print('Added updateFields to settings.xml')
    else:
        # Create settings.xml if not present
        os.makedirs(os.path.join(temp_dir, 'word'), exist_ok=True)
        with open(settings_xml_path, 'w', encoding='utf-8') as f:
            f.write('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n<w:settings xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:updateFields w:val="true"/></w:settings>')
        print('Created settings.xml with updateFields.')

    # Re-zip
    with zipfile.ZipFile(dst_path, 'w', zipfile.ZIP_DEFLATED) as zout:
        for root, dirs, files in os.walk(temp_dir):
            for file in files:
                full_path = os.path.join(root, file)
                rel_path = os.path.relpath(full_path, temp_dir)
                zout.write(full_path, rel_path)
    
    shutil.rmtree(temp_dir)
    print('Re-zipped docx with updateFields enabled.')

if __name__ == '__main__':
    generate_report()
