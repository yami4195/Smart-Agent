import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn
import zipfile
import tempfile
import shutil
import os

def update_entire_document():
    src_file = "Wegagen Bank Sc(software development)_original_backup.docx"
    dst_file = "Wegagen Bank Sc(software development).docx"
    
    doc = docx.Document(src_file)
    
    # Configure document styles
    styles = doc.styles
    
    # Title
    if 'Title' in styles:
        t_style = styles['Title']
        t_style.font.name = 'Arial'
        t_style.font.size = Pt(16)
        t_style.font.bold = True
        t_style.font.color.rgb = RGBColor(0x00, 0x20, 0x60)
    
    # Heading 1
    h1 = styles['Heading 1']
    h1.font.name = 'Arial'
    h1.font.size = Pt(15)
    h1.font.bold = True
    h1.font.color.rgb = RGBColor(0x00, 0x20, 0x60) # Deep Navy
    h1.paragraph_format.space_before = Pt(14)
    h1.paragraph_format.space_after = Pt(6)
    h1.paragraph_format.keep_with_next = True
    
    # Heading 2
    h2 = styles['Heading 2']
    h2.font.name = 'Arial'
    h2.font.size = Pt(12.5)
    h2.font.bold = True
    h2.font.color.rgb = RGBColor(0x1F, 0x4E, 0x78)
    h2.paragraph_format.space_before = Pt(10)
    h2.paragraph_format.space_after = Pt(4)
    h2.paragraph_format.keep_with_next = True
    
    # Heading 3
    h3 = styles['Heading 3']
    h3.font.name = 'Arial'
    h3.font.size = Pt(11)
    h3.font.bold = True
    h3.font.color.rgb = RGBColor(0x2E, 0x75, 0xB5)
    h3.paragraph_format.space_before = Pt(6)
    h3.paragraph_format.space_after = Pt(2)
    h3.paragraph_format.keep_with_next = True
    
    # Normal
    norm = styles['Normal']
    norm.font.name = 'Times New Roman'
    norm.font.size = Pt(12)
    norm.font.color.rgb = RGBColor(0x11, 0x11, 0x11)
    norm.paragraph_format.line_spacing = 1.15
    norm.paragraph_format.space_after = Pt(4)
    norm.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    # Clean Cover Page Text
    for p in doc.paragraphs[:15]:
        if 'FACULITY OF INFORMATICS' in p.text:
            p.text = 'FACULTY OF INFORMATICS'
            p.runs[0].font.name = 'Arial'
            p.runs[0].font.size = Pt(12)
            p.runs[0].font.bold = True
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        if 'Yeamlak sisay' in p.text:
            p.text = 'Student Name:  Yeamlak Sisay'
            p.runs[0].font.name = 'Times New Roman'
            p.runs[0].font.size = Pt(11)
            p.runs[0].font.bold = True
        if 'Hosting Company:  Wegagen Bank S.C' in p.text and not p.text.endswith('S.C.'):
            p.text = 'Hosting Company:  Wegagen Bank S.C.'
            p.runs[0].font.name = 'Times New Roman'
            p.runs[0].font.size = Pt(11)
            p.runs[0].font.bold = True
        if 'June 23, 2026 -September 1, 2026' in p.text:
            p.text = 'Duration of Apprenticeship:  June 23, 2026 - September 1, 2026'
            p.runs[0].font.name = 'Times New Roman'
            p.runs[0].font.size = Pt(11)
            p.runs[0].font.bold = True
        if 'Date of Submission:  September 18,2026' in p.text:
            p.text = '                                                              Date of Submission:  September 18, 2026'
            p.runs[0].font.name = 'Times New Roman'
            p.runs[0].font.size = Pt(11)
            p.runs[0].font.bold = True
            p.runs[0].font.italic = True

    # Preliminary section page breaks and formatting
    for p in doc.paragraphs:
        if p.style.name == 'Heading 1' and p.text.strip() in ['Declaration', 'Approval of the Mentor', 'Acknowledgements', 'Executive Summary', 'Chapter One: Background of Wegagen Bank', 'Chapter Two: My Practical Attachment Experience', 'Chapter Three: Benefits I Gained', 'Chapter Four: Conclusion and Recommendations', 'References']:
            p.paragraph_format.page_break_before = True
        
        # Approval text fix
        if 'software Development Internship at Wegagen Bank' in p.text:
            p.text = p.text.replace('software Development Internship at Wegagen Bank', 'Software Development Internship at Wegagen Bank')

        # Fix Figure 1.1 caption
        if 'Figure 1.1: Wegagen Bank Head Quarter' in p.text:
            p.text = 'Figure 1.1: Wegagen Bank Headquarters'
            p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(8)
            p.runs[0].font.name = 'Times New Roman'
            p.runs[0].font.italic = True
            p.runs[0].font.size = Pt(10)
            p.runs[0].font.bold = True

    # Format Chapter 1 Headings
    for p in doc.paragraphs:
        txt = p.text.strip()
        if txt == 'Introduction':
            p.text = '1.1 Introduction'
            p.style = 'Heading 2'
        elif txt == 'Brief History of Wegagen Bank':
            p.text = '1.2 Brief History of Wegagen Bank'
            p.style = 'Heading 2'
        elif txt == 'Main Products and Services of Wegagen Bank':
            p.text = '1.3 Main Products and Services of Wegagen Bank'
            p.style = 'Heading 2'
        elif txt == 'Deposit and Savings Services':
            p.text = '1.3.1 Deposit and Savings Services'
            p.style = 'Heading 3'
        elif txt == 'Credit and Loan Facilities':
            p.text = '1.3.2 Credit and Loan Facilities'
            p.style = 'Heading 3'
        elif txt == 'Digital and Electronic Banking (e-Banking)':
            p.text = '1.3.3 Digital and Electronic Banking (e-Banking)'
            p.style = 'Heading 3'
        elif txt == 'Interest-Free Banking (Wegagen Amanah)':
            p.text = '1.3.4 Interest-Free Banking (Wegagen Amanah)'
            p.style = 'Heading 3'
        elif txt == "Main Customers and End Users of Wegagen Bank's Products and Services":
            p.text = "1.4 Main Customers and End Users of Wegagen Bank's Products and Services"
            p.style = 'Heading 2'
        elif txt == '1. Retail and Individual Banking Customers':
            p.text = '1.4.1 Retail and Individual Banking Customers'
            p.style = 'Heading 3'
        elif txt == '2. Micro, Small, and Medium Enterprises (MSMEs)':
            p.text = '1.4.2 Micro, Small, and Medium Enterprises (MSMEs)'
            p.style = 'Heading 3'
        elif txt == '3. Corporate and Commercial Entities':
            p.text = '1.4.3 Corporate and Commercial Entities'
            p.style = 'Heading 3'
        elif txt.startswith('1.5') and 'Organization and Structure' in txt:
            p.text = '1.5 Organization and Structure'
            p.style = 'Heading 2'

    # Fix glued 1.6 in Chapter 1
    paras = list(doc.paragraphs)
    for i, p in enumerate(paras):
        if '1.6 Work Flow of Wegagen Bank (Core Application Management Services)' in p.text:
            p.text = p.text.replace('\n\n1.6 Work Flow of Wegagen Bank (Core Application Management Services)', '').replace('1.6 Work Flow of Wegagen Bank (Core Application Management Services)', '').strip()
            # Insert 1.6 Heading 2
            next_p = paras[i+1]
            p_16 = next_p.insert_paragraph_before('1.6 Workflow in Core Application Management Services', style='Heading 2')
            break

    # Format Chapter 2 Headings and Text
    paras = list(doc.paragraphs)
    for i, p in enumerate(paras):
        txt = p.text.strip()
        if txt == '2.1 How I Get into The Company':
            p.text = '2.1 How I Joined the Company'
            p.style = 'Heading 2'
        elif txt == 'Supervision and Technical Leadership':
            p.text = '2.1.2 Supervision and Technical Leadership'
            p.style = 'Heading 3'
        elif txt == '2.2 Workflow in the Core Application Management Service Department':
            p.style = 'Heading 2'
        elif 'During my internship, I group worked on a mobile application' in p.text:
            p.text = p.text.replace('I group worked on a mobile application related project', 'I worked on a mobile application project').replace('reviewed our progress', 'reviewed my progress')
        elif 'The internship provided us with an opportunity' in p.text:
            p.text = p.text.replace('provided us with an opportunity', 'provided me with an opportunity')
        elif 'This stage help me understand the project scope' in p.text:
            p.text = p.text.replace('help me understand', 'helped me understand')
        elif 'unique sequential digital token, logs the active queue state' in p.text:
            p.text = p.text.replace(
                'automated schema migrations. unique sequential digital token, logs the active queue state, and persists the transaction to queue history.',
                'automated schema migrations. When a user joins the queue, the system generates a unique sequential digital token, logs the active queue state, and persists the transaction to queue history.'
            )
        elif txt == '2.3.5 Mobile Application Development':
            p.text = '2.3.4 Mobile Application Development'
            p.style = 'Heading 3'
        elif txt == '2.3.6 Queue Management and Foreign-Exchange Rate Implementation':
            p.text = '2.3.5 Queue Management and Foreign-Exchange Rate Implementation'
            p.style = 'Heading 3'
        elif txt == '2.3.7 Quality Assurance, System Documentation, and Project Delivery':
            p.text = '2.3.6 Quality Assurance, System Documentation, and Project Delivery'
            p.style = 'Heading 3'
        elif txt == 'Development Environment and Materials':
            p.text = '2.4.1 Development Environment and Tools'
            p.style = 'Heading 3'
        elif txt == 'React Native':
            p.text = '2.4.2 Frontend and Mobile Technologies'
            p.style = 'Heading 3'
        elif txt == 'Node.js and Express.js':
            p.text = '2.4.3 Backend and Database Technologies'
            p.style = 'Heading 3'
        elif txt == 'Clerk Authentication':
            p.text = '2.4.4 Authentication, Real-Time Communication, and Testing'
            p.style = 'Heading 3'
        elif txt == 'Git and GitHub':
            p.text = '2.4.5 Version Control and Methodology'
            p.style = 'Heading 3'
        elif txt == '2.5 How well I performed My Work Tasks':
            p.text = '2.5 Evaluation of Work Performance'
            p.style = 'Heading 2'
        elif txt == '1. Learning React Native and Expo from the Ground Up':
            p.text = '2.6.1 Learning React Native and Expo from the Ground Up'
            p.style = 'Heading 3'
        elif txt == '2. Designing the Application Architecture from Scratch':
            # Distinguish challenges section from measures section
            if i < 225:
                p.text = '2.6.2 Designing the Application Architecture from Scratch'
            else:
                p.text = '2.7.2 Iterative Design and Incremental Architecture'
            p.style = 'Heading 3'
        elif txt == '3. The API Integration Bug':
            p.text = '2.6.3 API Integration and Client-Server Communication Bug'
            p.style = 'Heading 3'
        elif txt == '4. Balancing Learning with a Demanding Schedule':
            p.text = '2.6.4 Balancing Learning with a Demanding Commute and Schedule'
            p.style = 'Heading 3'
        elif txt == '1. Overcoming the React Native/Expo Learning Curve':
            p.text = '2.7.1 Overcoming the React Native and Expo Learning Curve'
            p.style = 'Heading 3'
        elif txt == '3. Fixing the API Integration Bug':
            p.text = '2.7.3 Systematic Debugging of API Endpoints'
            p.style = 'Heading 3'
        elif txt == '4. Managing the Demanding Schedule':
            p.text = '2.7.4 Time Management and Structured Routine'
            p.style = 'Heading 3'

    # Format Chapter 3 & 4 Headings and split multi-line paragraphs
    paras = list(doc.paragraphs)
    for i, p in enumerate(paras):
        txt = p.text.strip()
        if txt == 'Chapter Three: Benefits I Gained':
            p.text = 'Chapter Three: Benefits Gained'
        elif txt.startswith('1. Mobile Application Architecture\nWorking with'):
            parts = p.text.split('\n', 1)
            p.text = '3.2.1 Mobile Application Architecture'
            p.style = 'Heading 3'
            paras[i+1].insert_paragraph_before(parts[1], style='Normal')
        elif txt.startswith('2. API Design and Client-Server Communication\nThrough build'):
            parts = p.text.split('\n', 1)
            p.text = '3.2.2 API Design and Client-Server Communication'
            p.style = 'Heading 3'
            paras[i+1].insert_paragraph_before(parts[1], style='Normal')
        elif txt.startswith('3. Software Development in a Real-World, Production-Oriented Environment\nUnlike'):
            parts = p.text.split('\n', 1)
            p.text = '3.2.3 Software Development in a Real-World, Production-Oriented Environment'
            p.style = 'Heading 3'
            paras[i+1].insert_paragraph_before(parts[1], style='Normal')
        elif txt.startswith('Structured Onboarding for Interns\nOne of the'):
            parts = p.text.split('\n', 1)
            p.text = '4.2.1 Structured Onboarding for Interns'
            p.style = 'Heading 3'
            paras[i+1].insert_paragraph_before(parts[1], style='Normal')
        elif txt.startswith('Pre-Internship Skill Preparation\nSeveral of the'):
            parts = p.text.split('\n', 1)
            p.text = '4.2.2 Pre-Internship Skill Preparation'
            p.style = 'Heading 3'
            paras[i+1].insert_paragraph_before(parts[1], style='Normal')
        elif txt.startswith('University and Industry Collaboration\nThere is'):
            parts = p.text.split('\n', 1)
            p.text = '4.2.3 University and Industry Collaboration'
            p.style = 'Heading 3'
            paras[i+1].insert_paragraph_before(parts[1], style='Normal')

    # Add Figure Captions for Chapter 2 images
    paras = list(doc.paragraphs)
    for i, p in enumerate(paras):
        blips = p._p.xpath('.//a:blip/@r:embed')
        if 'rId10' in blips or 'rId11' in blips:
            p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_cap = paras[i+1].insert_paragraph_before('Figure 2.1: Smart Agent Mobile App – Home and Branch Discovery Screens')
            p_cap.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_cap.paragraph_format.space_before = Pt(4)
            p_cap.paragraph_format.space_after = Pt(8)
            p_cap.runs[0].font.name = 'Times New Roman'
            p_cap.runs[0].font.italic = True
            p_cap.runs[0].font.size = Pt(10)
            p_cap.runs[0].font.bold = True
        elif 'rId12' in blips:
            p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_cap = paras[i+1].insert_paragraph_before('Figure 2.2: Dedicated Teller and Employee Queue Management Interface')
            p_cap.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_cap.paragraph_format.space_before = Pt(4)
            p_cap.paragraph_format.space_after = Pt(8)
            p_cap.runs[0].font.name = 'Times New Roman'
            p_cap.runs[0].font.italic = True
            p_cap.runs[0].font.size = Pt(10)
            p_cap.runs[0].font.bold = True
        elif 'rId14' in blips:
            p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_cap = paras[i+1].insert_paragraph_before('Figure 2.3: Backend REST API Swagger/OpenAPI Documentation')
            p_cap.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_cap.paragraph_format.space_before = Pt(4)
            p_cap.paragraph_format.space_after = Pt(8)
            p_cap.runs[0].font.name = 'Times New Roman'
            p_cap.runs[0].font.italic = True
            p_cap.runs[0].font.size = Pt(10)
            p_cap.runs[0].font.bold = True

    # Locate Chapter One start paragraph to insert Table of Contents and List of Figures before it
    ch1_p = None
    for p in doc.paragraphs:
        if p.text.strip() == 'Chapter One: Background of Wegagen Bank':
            ch1_p = p
            break
            
    if ch1_p:
        def add_toc_item_before(target_p, title, page_str, level=1):
            p = target_p.insert_paragraph_before()
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.15
            p.paragraph_format.tab_stops.add_tab_stop(Inches(6.5), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
            
            if level == 1:
                p.paragraph_format.left_indent = Inches(0)
                r1 = p.add_run(title)
                r1.font.name = 'Arial'
                r1.bold = True
                r1.font.size = Pt(10.5)
                r1.font.color.rgb = RGBColor(0x00, 0x20, 0x60)
                r2 = p.add_run(f'\t{page_str}')
                r2.font.name = 'Times New Roman'
                r2.bold = True
                r2.font.size = Pt(11)
            elif level == 2:
                p.paragraph_format.left_indent = Inches(0.25)
                r1 = p.add_run(title)
                r1.font.name = 'Times New Roman'
                r1.bold = False
                r1.font.size = Pt(11)
                r2 = p.add_run(f'\t{page_str}')
                r2.font.name = 'Times New Roman'
                r2.bold = False
                r2.font.size = Pt(11)
            elif level == 3:
                p.paragraph_format.left_indent = Inches(0.5)
                r1 = p.add_run(title)
                r1.font.name = 'Times New Roman'
                r1.bold = False
                r1.font.size = Pt(10.5)
                r2 = p.add_run(f'\t{page_str}')
                r2.font.name = 'Times New Roman'
                r2.bold = False
                r2.font.size = Pt(10.5)

        # Table of Contents
        p_toc_head = ch1_p.insert_paragraph_before('Table of Contents', style='Heading 1')
        p_toc_head.paragraph_format.page_break_before = True
        
        # Word Dynamic TOC field
        fld_run = p_toc_head.add_run()
        r_toc = fld_run._r
        fldChar1 = parse_xml(r'<w:fldChar %s w:fldCharType="begin"/>' % nsdecls('w'))
        instrText = parse_xml(r'<w:instrText %s xml:space="preserve"> TOC \o "1-3" \h \z \u </w:instrText>' % nsdecls('w'))
        fldChar2 = parse_xml(r'<w:fldChar %s w:fldCharType="separate"/>' % nsdecls('w'))
        fldChar3 = parse_xml(r'<w:fldChar %s w:fldCharType="end"/>' % nsdecls('w'))
        r_toc.append(fldChar1)
        r_toc.append(instrText)
        r_toc.append(fldChar2)
        r_toc.append(fldChar3)

        add_toc_item_before(ch1_p, "Declaration", "ii", level=1)
        add_toc_item_before(ch1_p, "Approval of the Mentor", "iii", level=1)
        add_toc_item_before(ch1_p, "Acknowledgements", "iv", level=1)
        add_toc_item_before(ch1_p, "Executive Summary", "v", level=1)
        add_toc_item_before(ch1_p, "List of Figures", "vi", level=1)
        
        add_toc_item_before(ch1_p, "CHAPTER ONE: BACKGROUND OF WEGAGEN BANK", "1", level=1)
        add_toc_item_before(ch1_p, "1.1 Introduction", "1", level=2)
        add_toc_item_before(ch1_p, "1.2 Brief History of Wegagen Bank", "1", level=2)
        add_toc_item_before(ch1_p, "1.3 Main Products and Services of Wegagen Bank", "2", level=2)
        add_toc_item_before(ch1_p, "1.3.1 Deposit and Savings Services", "2", level=3)
        add_toc_item_before(ch1_p, "1.3.2 Credit and Loan Facilities", "2", level=3)
        add_toc_item_before(ch1_p, "1.3.3 Digital and Electronic Banking (e-Banking)", "3", level=3)
        add_toc_item_before(ch1_p, "1.3.4 Interest-Free Banking (Wegagen Amanah)", "3", level=3)
        add_toc_item_before(ch1_p, "1.4 Main Customers and End Users of Wegagen Bank's Products and Services", "3", level=2)
        add_toc_item_before(ch1_p, "1.4.1 Retail and Individual Banking Customers", "3", level=3)
        add_toc_item_before(ch1_p, "1.4.2 Micro, Small, and Medium Enterprises (MSMEs)", "4", level=3)
        add_toc_item_before(ch1_p, "1.4.3 Corporate and Commercial Entities", "4", level=3)
        add_toc_item_before(ch1_p, "1.5 Organization and Structure", "5", level=2)
        add_toc_item_before(ch1_p, "1.6 Workflow in Core Application Management Services", "6", level=2)

        add_toc_item_before(ch1_p, "CHAPTER TWO: PRACTICAL ATTACHMENT EXPERIENCE", "7", level=1)
        add_toc_item_before(ch1_p, "2.1 How I Joined the Company", "7", level=2)
        add_toc_item_before(ch1_p, "2.1.1 Placement within Wegagen Bank: Core Application Management Service", "7", level=3)
        add_toc_item_before(ch1_p, "2.1.2 Supervision and Technical Leadership", "8", level=3)
        add_toc_item_before(ch1_p, "2.2 Workflow in the Core Application Management Service Department", "8", level=2)
        add_toc_item_before(ch1_p, "2.3 Work Tasks Executed During the Practical Attachment", "9", level=2)
        add_toc_item_before(ch1_p, "2.3.1 Problem Identification and Requirement Analysis", "9", level=3)
        add_toc_item_before(ch1_p, "2.3.2 System Planning and Design", "9", level=3)
        add_toc_item_before(ch1_p, "2.3.3 Backend Architecture and Database Implementation", "10", level=3)
        add_toc_item_before(ch1_p, "2.3.4 Mobile Application Development", "11", level=3)
        add_toc_item_before(ch1_p, "2.3.5 Queue Management and Foreign-Exchange Rate Implementation", "12", level=3)
        add_toc_item_before(ch1_p, "2.3.6 Quality Assurance, System Documentation, and Project Delivery", "13", level=3)
        add_toc_item_before(ch1_p, "2.4 Materials and Methodology", "14", level=2)
        add_toc_item_before(ch1_p, "2.4.1 Development Environment and Tools", "14", level=3)
        add_toc_item_before(ch1_p, "2.4.2 Frontend and Mobile Technologies", "14", level=3)
        add_toc_item_before(ch1_p, "2.4.3 Backend and Database Technologies", "15", level=3)
        add_toc_item_before(ch1_p, "2.4.4 Authentication, Real-Time Communication, and Testing", "15", level=3)
        add_toc_item_before(ch1_p, "2.4.5 Version Control and Methodology", "16", level=3)
        add_toc_item_before(ch1_p, "2.5 Evaluation of Work Performance", "17", level=2)
        add_toc_item_before(ch1_p, "2.6 Challenges Faced During the Practical Attachment", "18", level=2)
        add_toc_item_before(ch1_p, "2.6.1 Learning React Native and Expo from the Ground Up", "18", level=3)
        add_toc_item_before(ch1_p, "2.6.2 Designing the Application Architecture from Scratch", "18", level=3)
        add_toc_item_before(ch1_p, "2.6.3 API Integration and Client-Server Communication Bug", "19", level=3)
        add_toc_item_before(ch1_p, "2.6.4 Balancing Learning with a Demanding Commute and Schedule", "19", level=3)
        add_toc_item_before(ch1_p, "2.7 Measures Taken to Overcome Challenges", "19", level=2)
        add_toc_item_before(ch1_p, "2.7.1 Overcoming the React Native and Expo Learning Curve", "19", level=3)
        add_toc_item_before(ch1_p, "2.7.2 Iterative Design and Incremental Architecture", "20", level=3)
        add_toc_item_before(ch1_p, "2.7.3 Systematic Debugging of API Endpoints", "20", level=3)
        add_toc_item_before(ch1_p, "2.7.4 Time Management and Structured Routine", "20", level=3)

        add_toc_item_before(ch1_p, "CHAPTER THREE: BENEFITS GAINED", "21", level=1)
        add_toc_item_before(ch1_p, "3.1 Practical Skills", "21", level=2)
        add_toc_item_before(ch1_p, "3.2 Theoretical Understanding Gained", "22", level=2)
        add_toc_item_before(ch1_p, "3.2.1 Mobile Application Architecture", "22", level=3)
        add_toc_item_before(ch1_p, "3.2.2 API Design and Client-Server Communication", "22", level=3)
        add_toc_item_before(ch1_p, "3.2.3 Software Development in a Real-World, Production-Oriented Environment", "23", level=3)
        add_toc_item_before(ch1_p, "3.3 Interpersonal and Communication Skills Gained", "23", level=2)
        add_toc_item_before(ch1_p, "3.4 Teamwork Skills", "24", level=2)
        add_toc_item_before(ch1_p, "3.5 Understanding of Work Ethics", "24", level=2)

        add_toc_item_before(ch1_p, "CHAPTER FOUR: CONCLUSION AND RECOMMENDATIONS", "25", level=1)
        add_toc_item_before(ch1_p, "4.1 Conclusion", "25", level=2)
        add_toc_item_before(ch1_p, "4.2 Recommendations", "26", level=2)
        add_toc_item_before(ch1_p, "4.2.1 Structured Onboarding for Interns", "26", level=3)
        add_toc_item_before(ch1_p, "4.2.2 Pre-Internship Skill Preparation", "26", level=3)
        add_toc_item_before(ch1_p, "4.2.3 University and Industry Collaboration", "27", level=3)

        add_toc_item_before(ch1_p, "REFERENCES", "28", level=1)

        # List of Figures
        p_lof_head = ch1_p.insert_paragraph_before('List of Figures', style='Heading 1')
        p_lof_head.paragraph_format.page_break_before = True
        add_toc_item_before(ch1_p, "Figure 1.1: Wegagen Bank Headquarters", "2", level=2)
        add_toc_item_before(ch1_p, "Figure 2.1: Smart Agent Mobile App – Home and Branch Discovery Screens", "11", level=2)
        add_toc_item_before(ch1_p, "Figure 2.2: Dedicated Teller and Employee Queue Management Interface", "12", level=2)
        add_toc_item_before(ch1_p, "Figure 2.3: Backend REST API Swagger/OpenAPI Documentation", "13", level=2)

    # Save
    doc.save(dst_file)
    print(f"Document saved to {dst_file}")

    # Update settings.xml for automatic field updating
    temp_dir = tempfile.mkdtemp()
    with zipfile.ZipFile(dst_file, 'r') as zin:
        zin.extractall(temp_dir)
    
    settings_xml_path = os.path.join(temp_dir, 'word', 'settings.xml')
    if os.path.exists(settings_xml_path):
        with open(settings_xml_path, 'r', encoding='utf-8') as f:
            content = f.read()
        if '<w:updateFields' not in content:
            content = content.replace('</w:settings>', '<w:updateFields w:val="true"/></w:settings>')
            with open(settings_xml_path, 'w', encoding='utf-8') as f:
                f.write(content)
            print('Added updateFields to settings.xml')
    
    with zipfile.ZipFile(dst_file, 'w', zipfile.ZIP_DEFLATED) as zout:
        for root, dirs, files in os.walk(temp_dir):
            for file in files:
                full_path = os.path.join(root, file)
                rel_path = os.path.relpath(full_path, temp_dir)
                zout.write(full_path, rel_path)
    
    shutil.rmtree(temp_dir)
    print('Re-zipped docx with updateFields enabled.')

if __name__ == '__main__':
    update_entire_document()
