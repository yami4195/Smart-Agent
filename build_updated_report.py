import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn
import zipfile
import shutil
import os

def update_document():
    src_file = "Wegagen Bank Sc(software development)_original_backup.docx"
    dst_file = "Wegagen Bank Sc(software development).docx"
    
    doc = docx.Document(src_file)
    
    # 1. Update Document Styles for professional look
    styles = doc.styles
    
    # Heading 1
    h1 = styles['Heading 1']
    h1.font.name = 'Calibri'
    h1.font.size = Pt(18)
    h1.font.bold = True
    h1.font.color.rgb = RGBColor(0x1B, 0x36, 0x5D)
    
    # Heading 2
    h2 = styles['Heading 2']
    h2.font.name = 'Calibri'
    h2.font.size = Pt(13.5)
    h2.font.bold = True
    h2.font.color.rgb = RGBColor(0x2F, 0x54, 0x96)
    
    # Heading 3
    h3 = styles['Heading 3']
    h3.font.name = 'Calibri'
    h3.font.size = Pt(11.5)
    h3.font.bold = True
    h3.font.color.rgb = RGBColor(0x1F, 0x38, 0x64)
    
    # Normal
    norm = styles['Normal']
    norm.font.name = 'Calibri'
    norm.font.size = Pt(11)
    norm.font.color.rgb = RGBColor(0x22, 0x22, 0x22)
    norm.paragraph_format.line_spacing = 1.15
    norm.paragraph_format.space_after = Pt(6)

    # Inspect all paragraphs
    print(f"Loaded original document with {len(doc.paragraphs)} paragraphs.")
    
    # Let's clean up cover page text
    for p in doc.paragraphs[:15]:
        if 'FACULITY OF INFORMATICS' in p.text:
            p.text = p.text.replace('FACULITY OF INFORMATICS', 'FACULTY OF INFORMATICS')
        if 'Yeamlak sisay' in p.text:
            p.text = p.text.replace('Yeamlak sisay', 'Yeamlak Sisay')
        if 'Wegagen Bank S.C' in p.text and not p.text.endswith('S.C.'):
            p.text = p.text.replace('Wegagen Bank S.C', 'Wegagen Bank S.C.')
        if 'June 23, 2026 -September 1, 2026' in p.text:
            p.text = p.text.replace('June 23, 2026 -September 1, 2026', 'June 23, 2026 - September 1, 2026')
        if 'September 18,2026' in p.text:
            p.text = p.text.replace('September 18,2026', 'September 18, 2026')

    # Approval of mentor text
    for p in doc.paragraphs[20:23]:
        if 'software Development Internship at Wegagen Bank' in p.text:
            p.text = p.text.replace('software Development Internship at Wegagen Bank', 'Software Development Internship at Wegagen Bank')

    # Figure 1.1 caption
    for p in doc.paragraphs:
        if 'Figure 1.1: Wegagen Bank Head Quarter' in p.text:
            p.text = 'Figure 1.1: Wegagen Bank Headquarters'
            p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.runs[0].font.italic = True
            p.runs[0].font.size = Pt(10)
            p.runs[0].font.bold = True

    # Fix broken sentence in Backend Architecture (2.3.3)
    for p in doc.paragraphs:
        if 'unique sequential digital token, logs the active queue state' in p.text:
            p.text = p.text.replace(
                'automated schema migrations. unique sequential digital token, logs the active queue state, and persists the transaction to queue history.',
                'automated schema migrations. When a user joins the queue, the system generates a unique sequential digital token, logs the active queue state, and persists the transaction to queue history.'
            )

    print("Basic text replacements done.")

update_document()
